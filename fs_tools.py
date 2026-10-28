"""
fs_tools.py
===========

File-system tools designed to be called by a Large Language Model (LLM).

Every tool returns plain, JSON-serialisable Python data (dicts / lists), so the
results can be sent straight back to the model. Tools never raise on expected
failures (missing file, unsupported type, corrupt PDF ...). Instead they return
a structured error, so the LLM can read it and recover.

Tools
-----
read_file(filepath)                    -> dict
list_files(directory, extension=None)  -> list
write_file(filepath, content)          -> dict
search_in_file(filepath, keyword)      -> dict

Supported read formats: .txt, .md, .pdf (pypdf), .docx (python-docx)
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Optional

SUPPORTED_READ_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
# write_file writes plain text only, so refuse binary formats that would be corrupted
BINARY_WRITE_BLOCKLIST = {".pdf", ".docx", ".doc", ".xlsx", ".png", ".jpg", ".zip"}


# --------------------------------------------------------------------------- #
# Internal helpers
# --------------------------------------------------------------------------- #
def _error(message: str, **extra) -> dict:
    """Build a consistent error payload."""
    return {"success": False, "error": message, **extra}


def _human_size(num_bytes: int) -> str:
    """Convert a byte count into a readable string, e.g. 2048 -> '2.0 KB'."""
    size = float(num_bytes)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024:
            return f"{int(size)} B" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024
    return f"{size:.1f} TB"


def _file_metadata(path: Path) -> dict:
    """Return basic metadata for a file."""
    stat = path.stat()
    return {
        "name": path.name,
        "path": str(path),
        "extension": path.suffix.lower(),
        "size_bytes": stat.st_size,
        "size_human": _human_size(stat.st_size),
        "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
    }


def _normalise_extension(extension: Optional[str]) -> Optional[str]:
    """'pdf', '.PDF', ' .pdf ' -> '.pdf'.  None / '' -> None."""
    if not extension or not extension.strip():
        return None
    ext = extension.strip().lower()
    return ext if ext.startswith(".") else f".{ext}"


def _extract_text(path: Path) -> tuple[str, dict]:
    """
    Extract text from a supported file.

    Returns (text, extra_metadata). Raises on failure; callers convert
    exceptions into structured errors.
    """
    ext = path.suffix.lower()
    extra: dict = {}

    if ext in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="replace")

    elif ext == ".pdf":
        from pypdf import PdfReader  # imported lazily so .txt works without pypdf

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ValueError("PDF is password-protected")
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(pages)
        extra["page_count"] = len(reader.pages)

    elif ext == ".docx":
        import docx  # python-docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        # Resumes often put skills in tables, so include table text too
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
        extra["paragraph_count"] = len(document.paragraphs)
        extra["table_count"] = len(document.tables)

    else:
        raise ValueError(f"Unsupported file type '{ext}'")

    return text.strip(), extra


# --------------------------------------------------------------------------- #
# Tool 1: read_file
# --------------------------------------------------------------------------- #
def read_file(filepath: str) -> dict:
    """
    Read a resume/document and extract its text.

    Args:
        filepath: Path to a .pdf, .docx, .txt or .md file.

    Returns:
        On success:
            {
              "success": True,
              "filepath": str,
              "content": str,             # extracted text
              "metadata": {name, extension, size_bytes, size_human, modified,
                           char_count, word_count, [page_count | paragraph_count]}
              "warning": str              # only if no text could be extracted
            }
        On failure:
            {"success": False, "error": str, "filepath": str}
    """
    if not filepath or not str(filepath).strip():
        return _error("filepath must be a non-empty string", filepath=filepath)

    path = Path(filepath).expanduser()

    if not path.exists():
        return _error(f"File not found: {filepath}", filepath=filepath)
    if not path.is_file():
        return _error(f"Path is not a file: {filepath}", filepath=filepath)

    ext = path.suffix.lower()
    if ext not in SUPPORTED_READ_EXTENSIONS:
        return _error(
            f"Unsupported file type '{ext or '(none)'}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_READ_EXTENSIONS))}",
            filepath=filepath,
        )

    try:
        text, extra = _extract_text(path)
    except ImportError as exc:
        return _error(
            f"Missing dependency ({exc}). Run: pip install -r requirements.txt",
            filepath=filepath,
        )
    except PermissionError:
        return _error(f"Permission denied: {filepath}", filepath=filepath)
    except Exception as exc:  # corrupt file, bad encoding, etc.
        return _error(
            f"Could not read '{path.name}': {type(exc).__name__}: {exc}",
            filepath=filepath,
        )

    metadata = _file_metadata(path)
    metadata.update(extra)
    metadata["char_count"] = len(text)
    metadata["word_count"] = len(text.split())

    result = {"success": True, "filepath": str(path), "content": text, "metadata": metadata}
    if not text:
        result["warning"] = (
            "No extractable text. The file may be empty or a scanned (image-only) PDF."
        )
    return result


# --------------------------------------------------------------------------- #
# Tool 2: list_files
# --------------------------------------------------------------------------- #
def list_files(directory: str, extension: Optional[str] = None) -> list:
    """
    List files in a directory (non-recursive), optionally filtered by extension.

    Args:
        directory: Directory to list.
        extension: Optional filter such as ".pdf", "pdf" or ".TXT" (case-insensitive).

    Returns:
        A list of metadata dicts sorted by name:
            [{"name", "path", "extension", "size_bytes", "size_human", "modified"}, ...]
        If the directory is invalid, a single-item list containing an error:
            [{"success": False, "error": "..."}]
        An empty list means the directory exists but no files matched.
    """
    if not directory or not str(directory).strip():
        return [_error("directory must be a non-empty string")]

    path = Path(directory).expanduser()
    if not path.exists():
        return [_error(f"Directory not found: {directory}")]
    if not path.is_dir():
        return [_error(f"Path is not a directory: {directory}")]

    ext_filter = _normalise_extension(extension)

    try:
        files = [
            _file_metadata(entry)
            for entry in path.iterdir()
            if entry.is_file()
            and not entry.name.startswith(".")  # skip hidden files like .DS_Store
            and (ext_filter is None or entry.suffix.lower() == ext_filter)
        ]
    except PermissionError:
        return [_error(f"Permission denied: {directory}")]

    return sorted(files, key=lambda f: f["name"].lower())


# --------------------------------------------------------------------------- #
# Tool 3: write_file
# --------------------------------------------------------------------------- #
def write_file(filepath: str, content: str, append: bool = False) -> dict:
    """
    Write text content to a file, creating parent directories if needed.

    Args:
        filepath: Destination path (e.g. "summaries/john_doe_summary.md").
        content:  Text to write (UTF-8).
        append:   If True, append instead of overwriting.

    Returns:
        On success:
            {"success": True, "filepath": str, "action": "created"|"overwritten"|"appended",
             "bytes_written": int, "created_directories": bool}
        On failure:
            {"success": False, "error": str, "filepath": str}
    """
    if not filepath or not str(filepath).strip():
        return _error("filepath must be a non-empty string", filepath=filepath)
    if content is None:
        return _error("content must not be None", filepath=filepath)

    path = Path(filepath).expanduser()

    if path.suffix.lower() in BINARY_WRITE_BLOCKLIST:
        return _error(
            f"write_file writes plain text and cannot produce '{path.suffix}' files. "
            "Use .txt or .md instead.",
            filepath=filepath,
        )
    if path.exists() and path.is_dir():
        return _error(f"Path is a directory: {filepath}", filepath=filepath)

    existed = path.exists()
    parent_missing = not path.parent.exists()

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        data = str(content)
        with path.open("a" if append else "w", encoding="utf-8") as fh:
            fh.write(data)
    except PermissionError:
        return _error(f"Permission denied: {filepath}", filepath=filepath)
    except OSError as exc:
        return _error(f"Could not write file: {exc}", filepath=filepath)

    if append and existed:
        action = "appended"
    elif existed:
        action = "overwritten"
    else:
        action = "created"

    return {
        "success": True,
        "filepath": str(path),
        "action": action,
        "bytes_written": len(data.encode("utf-8")),
        "created_directories": parent_missing,
    }


# --------------------------------------------------------------------------- #
# Tool 4: search_in_file
# --------------------------------------------------------------------------- #
def search_in_file(
    filepath: str,
    keyword: str,
    context_chars: int = 80,
    max_matches: int = 25,
) -> dict:
    """
    Case-insensitive keyword/phrase search inside a document.

    Multi-word phrases match across any whitespace, so "machine learning" still
    matches when a PDF line-break splits the words.

    Args:
        filepath:      File to search (.pdf, .docx, .txt, .md).
        keyword:       Word or phrase to find.
        context_chars: Characters of surrounding text to include on each side.
        max_matches:   Maximum matches returned (match_count is always the full total).

    Returns:
        {
          "success": True,
          "filepath": str,
          "keyword": str,
          "found": bool,
          "match_count": int,
          "matches": [{"line": int, "position": int, "match": str, "context": str}, ...],
          "truncated": bool
        }
        or {"success": False, "error": str, ...}
    """
    if not keyword or not str(keyword).strip():
        return _error("keyword must be a non-empty string", filepath=filepath, keyword=keyword)

    doc = read_file(filepath)
    if not doc["success"]:
        return _error(doc["error"], filepath=filepath, keyword=keyword)

    text = doc["content"]
    words = str(keyword).split()
    pattern = re.compile(r"\s+".join(re.escape(w) for w in words), re.IGNORECASE)

    matches = []
    total = 0
    for m in pattern.finditer(text):
        total += 1
        if len(matches) >= max_matches:
            continue
        start = max(0, m.start() - context_chars)
        end = min(len(text), m.end() + context_chars)
        snippet = " ".join(text[start:end].split())  # collapse newlines/whitespace
        context = ("..." if start > 0 else "") + snippet + ("..." if end < len(text) else "")
        matches.append(
            {
                "line": text.count("\n", 0, m.start()) + 1,
                "position": m.start(),
                "match": m.group(0),
                "context": context,
            }
        )

    return {
        "success": True,
        "filepath": doc["filepath"],
        "keyword": keyword,
        "found": total > 0,
        "match_count": total,
        "matches": matches,
        "truncated": total > len(matches),
    }


if __name__ == "__main__":
    # Quick manual check: python fs_tools.py
    import json

    print(json.dumps(list_files("resumes"), indent=2))