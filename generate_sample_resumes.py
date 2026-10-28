"""Generate dummy resumes in PDF, DOCX and TXT formats under ./resumes."""

from pathlib import Path

import docx
from fpdf import FPDF

OUT = Path("resumes")

RESUMES = {
    "resume_john_doe.pdf": """JOHN DOE
Senior Software Engineer | john.doe@example.com | Bengaluru, India

SUMMARY
Backend engineer with 6 years of experience building scalable APIs and ML pipelines.

SKILLS
Python, Django, FastAPI, PostgreSQL, Docker, Kubernetes, AWS, scikit-learn

EXPERIENCE
Senior Software Engineer, FinStack Technologies (2021 - Present)
- Designed Python microservices processing 2M transactions per day
- Led migration from monolith to FastAPI services, cutting latency by 40%
Software Engineer, DataWorks (2018 - 2021)
- Built ETL pipelines in Python and Airflow
- Deployed machine learning fraud models to production

EDUCATION
B.Tech Computer Science, IIT Kanpur, 2018""",

    "resume_priya_sharma.pdf": """PRIYA SHARMA
Java Developer | priya.sharma@example.com | Pune, India

SUMMARY
Java developer with 4 years of experience in enterprise banking applications.

SKILLS
Java, Spring Boot, Hibernate, MySQL, Kafka, Jenkins, Microservices

EXPERIENCE
Java Developer, CoreBank Solutions (2020 - Present)
- Built Spring Boot services for loan processing
- Integrated Kafka event streaming across 12 services

EDUCATION
B.E. Information Technology, Pune University, 2020""",

    "resume_alex_chen.docx": """ALEX CHEN
Data Analyst | alex.chen@example.com | Singapore

SUMMARY
Data analyst with 3 years of experience turning messy data into business insight.

SKILLS
SQL, Python (pandas, matplotlib), Tableau, Excel, A/B testing, Statistics

EXPERIENCE
Data Analyst, ShopRight Retail (2022 - Present)
- Automated weekly sales reporting with Python, saving 10 hours per week
- Built Tableau dashboards used by 50+ regional managers

EDUCATION
B.Sc. Statistics, National University of Singapore, 2021""",

    "resume_maria_garcia.docx": """MARIA GARCIA
Frontend Developer | maria.garcia@example.com | Madrid, Spain

SUMMARY
Frontend developer focused on accessible, high-performance web apps.

SKILLS
JavaScript, TypeScript, React, Next.js, Tailwind CSS, Jest, Figma

EXPERIENCE
Frontend Developer, Pixel Labs (2021 - Present)
- Rebuilt the customer portal in React and TypeScript
- Improved Lighthouse accessibility score from 68 to 97

EDUCATION
B.Sc. Software Engineering, Universidad Politecnica de Madrid, 2021""",

    "resume_rahul_verma.txt": """RAHUL VERMA
DevOps Engineer | rahul.verma@example.com | Lucknow, India

SUMMARY
DevOps engineer with 5 years automating cloud infrastructure.

SKILLS
AWS, Terraform, Kubernetes, Docker, GitHub Actions, Bash, Python scripting, Prometheus

EXPERIENCE
DevOps Engineer, CloudNine Systems (2020 - Present)
- Managed 40+ Kubernetes clusters on AWS EKS
- Wrote Python and Bash automation that reduced deployment time by 60%

EDUCATION
B.Tech Electronics, AKTU, 2019""",

    "resume_emily_watson.txt": """EMILY WATSON
Product Manager | emily.watson@example.com | London, UK

SUMMARY
Product manager with 7 years of experience in B2B SaaS.

SKILLS
Product strategy, Roadmapping, Agile/Scrum, JIRA, User research, SQL basics

EXPERIENCE
Senior Product Manager, Teamly (2019 - Present)
- Launched analytics module that grew ARR by 25%
- Ran 30+ customer discovery interviews per quarter

EDUCATION
MBA, London Business School, 2017""",
}


def make_pdf(path: Path, text: str) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 6, text)
    pdf.output(str(path))


def make_docx(path: Path, text: str) -> None:
    document = docx.Document()
    lines = text.splitlines()
    document.add_heading(lines[0], level=1)
    for line in lines[1:]:
        if line.isupper() and line.strip():
            document.add_heading(line.title(), level=2)
        else:
            document.add_paragraph(line)
    document.save(str(path))


def main() -> None:
    OUT.mkdir(exist_ok=True)
    for name, text in RESUMES.items():
        path = OUT / name
        if name.endswith(".pdf"):
            make_pdf(path, text)
        elif name.endswith(".docx"):
            make_docx(path, text)
        else:
            path.write_text(text, encoding="utf-8")
        print(f"created {path}")


if __name__ == "__main__":
    main()